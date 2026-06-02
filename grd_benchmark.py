"""
Honest GRD benchmark — scores the text-layer diffuser extractor against the
team's completed takeoffs. No hardcoded numbers, no fake model citation, and
precision is reported as N/A (not 100%) when we extract zero units.

Truth loading handles all three takeoff formats seen in the sample corpus:
  - standard : sheets ['TAKEOFF', 'TO', 'DATA']            → read 'DATA'
  - Haldeman : ['Triune Takeoff Haldeman', 'RawData']      → read 'RawData'
  - variants : ['TAKEOFF','DATA'] | ['Takeoff','RawData'] | ['TAKEOFF']
The GRD rows are the ones whose PRODUCT starts with 'AD' (air device).

A truth "unit" is a (mark, neck_size) pair carrying a Count/QTY. Recall counts
how many of the team's units we caught; precision how many of ours were real.
Matching is on (normalized mark, canonical neck) so '6X6' == 'rect:6x6'.

Usage
-----
  python grd_benchmark.py                                   # whole sample root
  python grd_benchmark.py --projects "Crunch" "Dick's"      # substring filter
  python grd_benchmark.py --root "<path>" --out grd_bench
  python grd_benchmark.py --marks-from truth                # use truth marks as
                                                            # the whitelist (debug;
                                                            # default is schedule)
"""
import argparse
import csv
import glob
import os
import sys
from collections import defaultdict
from pathlib import Path

from diffuser_extractor import (
    extract_diffuser_instances,
    normalize_neck_size,
)

DEFAULT_ROOT = r"C:/Users/JFL/Downloads/SAMPLE FILES 26.05.2026 EXTRACT/SAMPLE FILES 26.05.2026"

# Sheet preference: per-instance data sheets first, formatted sheet last.
_TRUTH_SHEET_PREF = ("DATA", "RAWDATA")
_HEADER_TOKENS = ("PRODUCT", "TAG", "NECK")


def _norm_mark(raw):
    if raw is None:
        return ""
    return str(raw).strip().upper()


def _norm_neck(raw):
    """Canonicalize a truth neck-size cell to the extractor's canon form."""
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s or s in (".", "-"):
        return ""
    canon = normalize_neck_size(s)
    return canon or ""


def _pick_truth_sheet(wb):
    upper = {s.upper(): s for s in wb.sheetnames}
    for pref in _TRUTH_SHEET_PREF:
        if pref in upper:
            return upper[pref]
    # Haldeman/standard formatted sheet — usable, header is still PRODUCT-based.
    for s in wb.sheetnames:
        if "TAKEOFF" in s.upper():
            return s
    return wb.sheetnames[0]


def load_truth_units(xlsx_path):
    """Return (units, n_rows) where units maps (mark, neck_canon) -> int count.

    Only PRODUCT rows starting with 'AD' (air devices = GRD) are counted.
    Returns ({}, 0) on any failure — caller treats that as 'no truth'.
    """
    import openpyxl
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception:
        return {}, 0

    sheet = _pick_truth_sheet(wb)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}, 0

    # Locate header row (first row carrying PRODUCT + TAG + NECK tokens).
    header_idx = None
    for i, row in enumerate(rows[:15]):
        cells = [str(c).upper() if c is not None else "" for c in row]
        if all(any(tok in c for c in cells) for tok in _HEADER_TOKENS):
            header_idx = i
            break
    if header_idx is None:
        return {}, 0

    header = [str(c).strip().upper() if c is not None else "" for c in rows[header_idx]]

    def col(*names):
        for nm in names:
            for ci, h in enumerate(header):
                if nm in h:
                    return ci
        return None

    c_prod = col("PRODUCT")
    c_tag = col("TAG")
    c_neck = col("NECK")
    c_qty = col("COUNT", "QTY")
    if None in (c_prod, c_tag, c_neck, c_qty):
        return {}, 0

    units = defaultdict(int)
    n_rows = 0
    for row in rows[header_idx + 1:]:
        if c_prod >= len(row):
            continue
        prod = row[c_prod]
        if prod is None or not str(prod).strip().upper().startswith("AD"):
            continue
        mark = _norm_mark(row[c_tag] if c_tag < len(row) else None)
        neck = _norm_neck(row[c_neck] if c_neck < len(row) else None)
        try:
            qty = int(float(row[c_qty])) if c_qty < len(row) and row[c_qty] not in (None, "") else 1
        except (ValueError, TypeError):
            qty = 1
        if not mark:
            continue
        units[(mark, neck)] += qty
        n_rows += 1
    return dict(units), n_rows


def _resolve_plan_pages(pdf_path):
    import fitz
    with fitz.open(str(pdf_path)) as doc:
        total = doc.page_count
    try:
        from takeoff_cli import find_mechanical_pages
        plan = find_mechanical_pages(str(pdf_path))
    except Exception:
        plan = []
    if not plan:
        plan = list(range(total))
    return plan, total


def _schedule_marks(pdf_path):
    try:
        from schedule_parser import parse_pdf_schedules
        _s, marks, _md, _l, _sm, _v = parse_pdf_schedules(str(pdf_path))
        return set(marks) if marks else None
    except Exception:
        return None


def our_units_from_instances(instances):
    units = defaultdict(int)
    for inst in instances:
        mark = _norm_mark(inst.get("mark"))
        neck = inst.get("neck_size_canon") or ""
        if not mark:
            continue
        units[(mark, neck)] += 1
    return dict(units)


def _collapse_to_mark(units):
    """Collapse a (mark, neck)->count map to mark->count (ignore neck size)."""
    out = defaultdict(int)
    for (mark, _neck), c in units.items():
        out[mark] += c
    return dict(out)


def score(truth, ours):
    """Score at two granularities.

    unit-level  : strict (mark, neck) match — what a real takeoff needs.
    mark-level  : (mark only) — did we detect the device at all, size aside?
                  The gap between the two is the 'found the mark but not the
                  neck size' bucket, which dominates bare-mark drawings.

    Returns a dict. precision is None when we extracted zero units (avoids the
    fake-100% bug the colleague's harness had).
    """
    def _one(t, o):
        t_tot, o_tot = sum(t.values()), sum(o.values())
        m = sum(min(t.get(k, 0), o.get(k, 0)) for k in set(t) | set(o))
        return ((m / t_tot) if t_tot else None,
                (m / o_tot) if o_tot else None, m, t_tot, o_tot)

    u_rec, u_prec, u_m, u_t, u_o = _one(truth, ours)
    mk_rec, mk_prec, mk_m, mk_t, mk_o = _one(_collapse_to_mark(truth), _collapse_to_mark(ours))
    return {
        "recall": u_rec, "precision": u_prec, "matched": u_m,
        "truth_total": u_t, "our_total": u_o,
        "mark_recall": mk_rec, "mark_precision": mk_prec, "mark_matched": mk_m,
    }


def find_plan_pdf(proj_dir):
    ps = os.path.join(proj_dir, "Plans_Specs")
    if not os.path.isdir(ps):
        return None
    pdfs = [p for p in glob.glob(os.path.join(ps, "*.pdf"))]
    # Prefer a non-RCP plan set (reflected-ceiling plans rarely carry the schedule).
    non_rcp = [p for p in pdfs if "RCP" not in os.path.basename(p).upper()]
    pool = non_rcp or pdfs
    # Pick the largest (usually the full drawing set, not a single schedule sheet).
    return max(pool, key=os.path.getsize) if pool else None


def find_truth_xlsx(proj_dir):
    ct = os.path.join(proj_dir, "Completed Takeoff")
    if not os.path.isdir(ct):
        return None
    xs = [f for f in glob.glob(os.path.join(ct, "*.xlsx"))
          if "takeoff" in os.path.basename(f).lower()]
    return xs[0] if xs else None


def fmt_pct(v):
    return "  N/A" if v is None else f"{v * 100:5.1f}%"


def run(root, projects_filter, out_dir, marks_source):
    root = Path(root)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    proj_dirs = sorted(
        d for d in (root / p for p in os.listdir(root))
        if d.is_dir() and os.path.isdir(d / "Plans_Specs")
    )
    if projects_filter:
        proj_dirs = [d for d in proj_dirs
                     if any(f.lower() in d.name.lower() for f in projects_filter)]

    results = []
    for d in proj_dirs:
        name = d.name
        pdf = find_plan_pdf(str(d))
        truth_xlsx = find_truth_xlsx(str(d))
        if not pdf or not truth_xlsx:
            print(f"  skip {name[:40]}: missing plan or truth")
            continue

        truth, truth_rows = load_truth_units(truth_xlsx)
        if not truth:
            print(f"  skip {name[:40]}: no GRD truth units loaded")
            results.append((name, "no_truth", None, None, 0, 0, 0, None, ""))
            continue

        plan_pages, total = _resolve_plan_pages(pdf)
        if marks_source == "truth":
            valid = {m for (m, _n) in truth} or None
        else:
            valid = _schedule_marks(pdf)

        try:
            instances, _warn = extract_diffuser_instances(
                str(pdf), plan_pages, valid_marks=valid, enable_ocr_supplement=False)
            status = "ok"
        except Exception as e:                      # noqa: BLE001
            print(f"  ERR  {name[:40]}: {e}")
            results.append((name, "error", None, None, 0, sum(truth.values()), 0, None, str(e)[:60]))
            continue

        ours = our_units_from_instances(instances)
        sc = score(truth, ours)
        style = "valid=%d" % (len(valid) if valid else 0)
        results.append((name, status, sc["recall"], sc["precision"], sc["matched"],
                        sc["truth_total"], sc["our_total"], sc["mark_recall"], style))
        print(f"  {name[:42]:44} unit {fmt_pct(sc['recall'])}  mark {fmt_pct(sc['mark_recall'])}"
              f"  prec {fmt_pct(sc['precision'])}"
              f"  ({sc['matched']}/{sc['truth_total']} truth, {sc['our_total']} ours)")

    # ── Write CSV ─────────────────────────────────────────────────────────────
    csv_path = out_dir / "grd_benchmark_results.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["project", "status", "unit_recall_pct", "precision_pct",
                    "matched_units", "truth_units", "our_units", "mark_recall_pct", "note"])
        for (name, status, rec, prec, m, t, o, mk, note) in results:
            w.writerow([name, status,
                        "" if rec is None else round(rec * 100, 1),
                        "" if prec is None else round(prec * 100, 1),
                        m, t, o,
                        "" if mk is None else round(mk * 100, 1), note])

    # ── Aggregate (only scored projects with truth) ──────────────────────────
    scored = [r for r in results if r[1] == "ok" and r[5] > 0]
    if scored:
        recalls = [r[2] for r in scored if r[2] is not None]
        mark_recalls = [r[7] for r in scored if r[7] is not None]
        tot_truth = sum(r[5] for r in scored)
        tot_match = sum(r[4] for r in scored)
        med = sorted(recalls)[len(recalls) // 2] if recalls else None
        med_mark = sorted(mark_recalls)[len(mark_recalls) // 2] if mark_recalls else None
        micro = tot_match / tot_truth if tot_truth else None
        print("\n" + "=" * 60)
        print(f"  scored projects: {len(scored)}  |  no_truth/err: {len(results) - len(scored)}")
        print(f"  micro UNIT recall (all units): {fmt_pct(micro)}  ({tot_match}/{tot_truth})")
        print(f"  median per-project UNIT recall: {fmt_pct(med)}")
        print(f"  median per-project MARK recall: {fmt_pct(med_mark)}")
        print("=" * 60)
    print(f"  → {csv_path}")
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(description="Honest GRD benchmark vs team takeoffs")
    ap.add_argument("--root", default=DEFAULT_ROOT)
    ap.add_argument("--projects", nargs="+", default=None, help="substring filter")
    ap.add_argument("--out", default="grd_bench_output")
    ap.add_argument("--marks-from", choices=["schedule", "truth"], default="schedule",
                    help="source of the valid-mark whitelist (default: schedule = production-correct)")
    args = ap.parse_args(argv)
    return run(args.root, args.projects, args.out, args.marks_from)


if __name__ == "__main__":
    raise SystemExit(main())
